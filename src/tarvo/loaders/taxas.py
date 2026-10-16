"""Planilha de taxas contratadas (mantida à mão pelo financeiro).

Particularidades tratadas:
- cabeçalho localizado pelo texto 'Lojista'/'CNPJ', não por posição;
- abas com 'nao usar' / 'antigo' no nome são ignoradas;
- taxa como fração (0.0149), como texto ('1,49%'), 'padrão' ou vazia;
- vigência em texto livre: 'até 14/09/2026', 'a partir de 15/09/2026', '01/01/2026';
- linha tachada = lojista inativo; comentário de célula é preservado;
- CNPJ com máscara, sem máscara ou numérico sem zero à esquerda.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

from ..util import cnpj, taxa_pct

# nome canônico da coluna -> trechos aceitos no cabeçalho (sem acento, minúsculo)
CABECALHOS = {
    "lojista": ["lojista"],
    "cnpj": ["cnpj"],
    "ec_orbe": ["ec orbe", "ec"],
    "debito": ["debito"],
    "credito_vista": ["credito a vista", "cred. a vista"],
    "credito_2_6": ["2-6"],
    "credito_7_12": ["7-12"],
    "pix": ["pix"],
    "antecipacao": ["antecipacao"],
    "spread_remessa": ["spread"],
    "vigencia": ["vigencia"],
    "obs": ["obs"],
}
TAXAS = ["debito", "credito_vista", "credito_2_6", "credito_7_12", "pix", "antecipacao", "spread_remessa"]

INICIO = date(1900, 1, 1)
FIM = date(2999, 12, 31)


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _mapear_cabecalho(row) -> dict[str, int] | None:
    textos = {i: _norm(c.value) for i, c in enumerate(row) if c.value is not None}
    if not any(t == "lojista" for t in textos.values()) or not any("cnpj" in t for t in textos.values()):
        return None
    idx = {}
    for canon, chaves in CABECALHOS.items():
        for i, t in textos.items():
            if i in idx.values():
                continue
            if any(t == k or (len(k) >= 3 and k in t) for k in chaves):
                idx[canon] = i
                break
    return idx


def _vigencia(v) -> tuple[date, date, str]:
    """Converte a célula de vigência em (início, fim, texto original)."""
    if v is None:
        return INICIO, FIM, ""
    if isinstance(v, (datetime, date)):
        d = v.date() if isinstance(v, datetime) else v
        return d, FIM, str(v)
    s = str(v)
    m = re.search(r"(\d{2}/\d{2}/\d{4})", s)
    if not m:
        return INICIO, FIM, s
    d = datetime.strptime(m.group(1), "%d/%m/%Y").date()
    ns = _norm(s)
    if ns.startswith("ate"):
        return INICIO, d, s
    return d, FIM, s  # 'a partir de' ou data simples = início de vigência


def carregar(path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, data_only=True)
    linhas = []
    for ws in wb.worksheets:
        if any(k in _norm(ws.title) for k in ("nao usar", "antigo")):
            continue
        idx = None
        for row in ws.iter_rows():
            if idx is None:
                idx = _mapear_cabecalho(row)
                continue
            get = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None
            nome_c = get("lojista")
            if nome_c is None or nome_c.value is None or str(nome_c.value).strip().startswith("*"):
                continue
            ini, fim, vig_txt = _vigencia(get("vigencia").value if get("vigencia") else None)
            obs = get("obs").value if get("obs") else None
            reg = {
                "aba": ws.title,
                "linha_planilha": nome_c.row,
                "ref_origem": f"{ws.title}!L{nome_c.row}",
                "lojista_planilha": str(nome_c.value).strip(),
                "cnpj": cnpj(get("cnpj").value if get("cnpj") else None),
                "ec_orbe": (str(get("ec_orbe").value).strip() if get("ec_orbe") and get("ec_orbe").value else None),
                "vig_inicio": ini,
                "vig_fim": fim,
                "vigencia_texto": vig_txt,
                "obs": obs,
                "comentario": nome_c.comment.text.strip() if nome_c.comment else None,
                "inativo": bool(nome_c.font and nome_c.font.strike) or "inativo" in _norm(obs or ""),
            }
            for t in TAXAS:
                c = get(t)
                pct, marcador = taxa_pct(c.value if c is not None else None)
                reg[t] = pct
                reg[f"{t}_marcador"] = marcador
            linhas.append(reg)
    df = pd.DataFrame(linhas)
    df["arquivo_origem"] = path.name
    return df


def regra_spread_padrao(path: Path) -> str | None:
    """Devolve a nota de rodapé sobre spread padrão, se existir (ex.: '* spread em branco = padrão')."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().startswith("*") and "spread" in _norm(c.value):
                    return c.value.strip()
    return None
